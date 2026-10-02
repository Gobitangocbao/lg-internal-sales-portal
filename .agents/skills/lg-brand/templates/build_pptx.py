#!/usr/bin/env python3
"""
build_pptx.py - build an LG-branded PowerPoint deck with a real LG theme.

Why this exists: the .pptx templates shipped in LG's asset pack carry the DEFAULT OFFICE THEME
(Calibri, accent #4472C4) with brand appearance pasted on top as pictures. Editing those keeps
the wrong theme, so every new text box inherits Calibri and Office blue. This script rewrites
the theme part - colour scheme and font scheme - so the deck is brand-correct at its root.

Usage as a library:

    from build_pptx import LGDeck
    d = LGDeck()                              # 16:9, 13.333 x 7.5 in
    d.title_slide("Channel plan 2026", "GTM Channel Development", slogan=True)
    d.section("Where we are")
    d.content_slide("Three things changed this quarter",
                    ["Coverage up in MT", "Display share flat", "Sell-out ahead of plan"])
    d.save("deck.pptx")

Or run it directly for a demo deck:

    python3 templates/build_pptx.py out.pptx

Requires python-pptx. Install the LG EI fonts on the machine first, and after opening the file
turn on font embedding (File > Options > Save > Embed fonts in the file) so the deck survives
being sent to someone else.
"""
import copy
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "..", "assets")

# --- palette (BI Guidelines V5.2 p.40; see references/color.md) ---
ACTIVE_RED = RGBColor(0xFD, 0x31, 0x2E)
LG_RED     = RGBColor(0xA5, 0x00, 0x34)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
BLACK      = RGBColor(0x00, 0x00, 0x00)
WG01       = RGBColor(0x26, 0x26, 0x26)
WG02       = RGBColor(0x4A, 0x49, 0x46)
WG03       = RGBColor(0x71, 0x6F, 0x6A)
WG04       = RGBColor(0xCB, 0xC8, 0xC2)
WG05       = RGBColor(0xE6, 0xE1, 0xD6)
WG06       = RGBColor(0xF0, 0xEC, 0xE4)
WG07       = RGBColor(0xF6, 0xF3, 0xEB)

HEADLINE = "LG EI Headline"
# LG EI Text weights each declare their own family name in the font file, and PowerPoint has no
# way to unify them. So the exact string per weight is what must be written, with bold OFF.
TEXT_REG = "LG EI Text Regular"
TEXT_SB  = "LG EI Text SemiBold"
TEXT_BD  = "LG EI Text Bold"

THEME_COLORS = [
    ("dk1", "262626"), ("lt1", "F6F3EB"), ("dk2", "4A4946"), ("lt2", "F0ECE4"),
    ("accent1", "FD312E"), ("accent2", "A50034"), ("accent3", "4A4946"),
    ("accent4", "716F6A"), ("accent5", "CBC8C2"), ("accent6", "E6E1D6"),
    ("hlink", "A50034"), ("folHlink", "716F6A"),
]

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _brand_theme(prs):
    """Rewrite the theme's colour and font scheme in place."""
    part = prs.slide_masters[0].part.part_related_by(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme")
    root = part._element if hasattr(part, "_element") else None
    if root is None:
        from lxml import etree
        root = etree.fromstring(part.blob)

    scheme = root.find(f"{A}themeElements/{A}clrScheme")
    if scheme is not None:
        for name, hexval in THEME_COLORS:
            el = scheme.find(f"{A}{name}")
            if el is None:
                continue
            for child in list(el):
                el.remove(child)
            from lxml import etree
            srgb = etree.SubElement(el, f"{A}srgbClr")
            srgb.set("val", hexval)

    fonts = root.find(f"{A}themeElements/{A}fontScheme")
    if fonts is not None:
        for tag, face in ((f"{A}majorFont", HEADLINE), (f"{A}minorFont", TEXT_REG)):
            grp = fonts.find(tag)
            if grp is None:
                continue
            for t in ("latin", "ea", "cs"):
                el = grp.find(f"{A}{t}")
                if el is not None:
                    el.set("typeface", face if t == "latin" else "")

    from lxml import etree
    # python-pptx's default template leaves Times New Roman / Arial in the latent
    # default text styles. Existing runs override them, but any NEW text box a person
    # adds later inherits them - so scrub them at the root.
    for el in root.iter(f"{A}latin"):
        if el.get("typeface", "").lower() in ("times new roman", "arial", "calibri",
                                              "calibri light", "+mj-lt", "+mn-lt"):
            el.set("typeface", HEADLINE if "major" in str(el.getparent().tag) else TEXT_REG)
    part._blob = etree.tostring(root, xml_declaration=True,
                                encoding="UTF-8", standalone=True)
    try:
        part.blob  # some python-pptx versions cache
    except Exception:
        pass
    part._element = root
    return part


class LGDeck:
    """A 16:9 deck built on the LG grid.

    Grid, per BI Guidelines p.76-78, on a 13.333 x 7.5 in slide:
        margin      = 5% of the shortest edge (7.5 in) = 0.375 in
        gutter      = margin / 2                       = 0.1875 in
        symbol h    = 1.3 x margin                     = 0.4875 in
        slogan h    = 0.9 x margin (sign-off)          = 0.3375 in

    1.3X and 0.9X measure the VISIBLE mark, not the file. The PNGs carry clear-space
    padding, so the placed picture is larger: logo file height = 1.3X / 0.668 = 0.730 in,
    slogan file height = 0.9X / 0.669 = 0.504 in. Placing the file at 1.3X directly - the
    obvious reading - lands the symbol about a third undersized.
    """

    def __init__(self, width_in=13.333, height_in=7.5):
        self.prs = Presentation()
        self.prs.slide_width = Inches(width_in)
        self.prs.slide_height = Inches(height_in)
        self.W, self.H = width_in, height_in
        self.margin = 0.05 * min(width_in, height_in)
        self.gutter = self.margin / 2
        self.symbol_h = 1.30 * self.margin          # the rule
        self.logo_h = self.symbol_h / 0.668         # compact LGE_Logo_* file height
        self.slogan_ink_h = 0.90 * self.margin      # the rule
        self.slogan_h = self.slogan_ink_h / 0.669   # slogan file height
        self.col = (self.W - 2 * self.margin - 8 * self.gutter) / 9
        self.row = (self.H - 2 * self.margin - 8 * self.gutter) / 9
        _brand_theme(self.prs)

    # ---------- primitives ----------

    def _blank(self, bg=WG07):
        layout = self.prs.slide_layouts[6]  # blank
        s = self.prs.slides.add_slide(layout)
        bgfill = s.background.fill
        bgfill.solid()
        bgfill.fore_color.rgb = bg
        return s

    def _logo(self, slide, dark_bg=False, position="upper-left"):
        # NOTE: there is no all-white "LGE_Logo_Mono_White" for the symbol+LG lockup in the
        # asset pack - only the 2D+LG-Electronics version has one. On dark backgrounds this
        # uses HeritageRed + White logotype, which p.21 approves for black and image grounds.
        name = ("LGE_Logo_HeritageRed_White_RGB.png" if dark_bg
                else "LGE_Logo_HeritageRed_Grey_RGB.png")
        path = os.path.join(ASSETS, "logo", name)
        # canvas aspect of LGE_Logo_* is 2081 x 1127; size by height, width follows
        h = self.logo_h
        self.logo_w = h * (2081 / 1127)
        if position == "upper-left":
            left, top = self.margin, self.margin
        elif position == "upper-right":
            left, top = self.W - self.margin - self.logo_w, self.margin
        elif position == "upper-center":
            left, top = (self.W - self.logo_w) / 2, self.margin
        elif position == "middle-left":
            left, top = self.margin, (self.H - h) / 2
        elif position == "bottom-left":
            left, top = self.margin, self.H - self.margin - h
        else:
            raise ValueError(
                f"{position!r} is not one of the five permitted logo positions: "
                "upper-left, upper-center, upper-right, middle-left, bottom-left")
        slide.shapes.add_picture(path, Inches(left), Inches(top), height=Inches(h))

    def _slogan(self, slide, white=False):
        """Sign-off slogan, bottom-right. Deliberately opposite the logo: the guideline
        forbids placing the logo and slogan next to each other."""
        name = ("LGE_Electronics_Slogan_Horizontal_Mono_White_RGB.png" if white
                else "LGE_Electronics_Slogan_Horizontal_Mono_ActiveRed_RGB.png")
        path = os.path.join(ASSETS, "slogan", name)
        w = self.slogan_h * (3832 / 870)
        slide.shapes.add_picture(
            path,
            Inches(self.W - self.margin - w),
            Inches(self.H - self.margin - self.slogan_h),
            height=Inches(self.slogan_h))

    def _text(self, slide, left, top, width, height, runs, align=PP_ALIGN.LEFT,
              anchor=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        for i, (txt, font, size, color, space_after) in enumerate(runs):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_after = Pt(space_after)
            r = p.add_run()
            r.text = txt
            r.font.name = font
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = False   # never faux-bold an already-bold LG EI Text family
            r.font.italic = False
        return box

    def _accent_rule(self, slide, left, top, width=None):
        from pptx.enum.shapes import MSO_SHAPE
        w = width or self.col
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top),
                                     Inches(w), Pt(4))
        shp.fill.solid()
        shp.fill.fore_color.rgb = ACTIVE_RED
        shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    # ---------- slide types ----------

    def title_slide(self, title, eyebrow=None, dark=False, slogan=True):
        s = self._blank(WG01 if dark else WG07)
        self._logo(s, dark_bg=dark)
        body_w = 6 * self.col + 5 * self.gutter
        top = self.margin + 2 * (self.row + self.gutter)
        self._accent_rule(s, self.margin, top)
        y = top + 0.16
        runs = []
        if eyebrow:
            runs.append((eyebrow.upper(), TEXT_SB, 13, ACTIVE_RED, 10))
        runs.append((title, HEADLINE, 44, WHITE if dark else WG01, 0))
        self._text(s, self.margin, y, body_w, self.row * 3, runs)
        if slogan:
            self._slogan(s, white=dark)
        return s

    def section(self, title, number=None, dark=True):
        s = self._blank(WG01 if dark else WG06)
        self._logo(s, dark_bg=dark)
        runs = []
        if number:
            runs.append((str(number), HEADLINE, 18, ACTIVE_RED, 8))
        runs.append((title, HEADLINE, 36, WHITE if dark else WG01, 0))
        self._text(s, self.margin, self.H / 2 - self.row,
                   6 * self.col + 5 * self.gutter, self.row * 2, runs,
                   anchor=MSO_ANCHOR.MIDDLE)
        return s

    def content_slide(self, title, bullets, dark=False):
        s = self._blank(WG01 if dark else WG07)
        self._logo(s, dark_bg=dark)
        fg = WHITE if dark else WG01
        body = WG04 if dark else WG02
        top = self.margin + self.row + self.gutter
        self._text(s, self.margin, top, 7 * self.col + 6 * self.gutter, self.row * 1.4,
                   [(title, HEADLINE, 28, fg, 0)])
        runs = [(f"—  {b}", TEXT_REG, 18, body, 12) for b in bullets]
        self._text(s, self.margin, top + self.row * 1.4,
                   6 * self.col + 5 * self.gutter, self.row * 5, runs)
        return s

    def save(self, path):
        self._scrub_default_fonts()
        self.prs.save(path)
        return path

    def _scrub_default_fonts(self):
        """python-pptx ships Times New Roman / Arial in presentation.xml's defaultTextStyle
        and in the notes master. Nothing visible uses them, but a text box added later in
        PowerPoint inherits them - which is exactly how a brand-correct deck drifts."""
        from lxml import etree
        for part in (self.prs.part,) + tuple(m.part for m in self.prs.slide_masters):
            root = part._element
            for el in root.iter(f"{A}latin"):
                tf = el.get("typeface", "").lower()
                if tf in ("times new roman", "arial", "calibri", "calibri light"):
                    el.set("typeface", TEXT_REG)
            for el in root.iter(f"{A}cs"):
                if el.get("typeface", "").lower() in ("times new roman", "arial"):
                    el.set("typeface", "")


def _demo(out):
    d = LGDeck()
    d.title_slide("Channel development plan", eyebrow="GTM  ·  Vietnam", slogan=True)
    d.section("Where we are", number="01")
    d.content_slide("Three things changed this quarter", [
        "Coverage expanded across modern trade",
        "Display share held flat against plan",
        "Sell-out running ahead of the quarterly target",
    ])
    d.save(out)
    print(f"wrote {out}")
    print("Next: python3 scripts/brand_check.py", out)
    print("Then open it and turn on font embedding before sharing.")


if __name__ == "__main__":
    _demo(sys.argv[1] if len(sys.argv) > 1 else "lg-deck.pptx")
